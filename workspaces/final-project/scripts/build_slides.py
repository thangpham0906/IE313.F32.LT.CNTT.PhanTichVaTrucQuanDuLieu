"""Tạo slide Nhóm 10 từ template môn học, dùng kết quả đã kiểm chứng."""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
REPORTS = ROOT / "reports"
NAVY = "15344C"
TEAL = "087E8B"
MUTED = "536777"
WHITE = "FFFFFF"
PALE = "EEF5F7"
GOLD = "D49019"


def main():
    """Giữ khổ 16:9 và theme gốc, tùy chỉnh nội dung theo chỉ dẫn template."""
    deck = Presentation(ROOT / "docs/Slides.pptx")
    for item in list(deck.slides._sldIdLst):
        deck.part.drop_rel(item.rId)
        deck.slides._sldIdLst.remove(item)
    layout = next(layout for layout in deck.slide_layouts if layout.name == "Blank")

    def box(slide, x, y, w, h, color):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(color)
        shape.line.fill.background()
        return shape

    def text(slide, value, x, y, w, h, size=22, color=NAVY, bold=False):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = shape.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Inches(0.02)
        frame.margin_top = frame.margin_bottom = Inches(0.02)
        for i, line in enumerate(value.split("\n")):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.space_after = Pt(10)
            paragraph.font.name = "DejaVu Sans"
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(color)
        return shape

    def picture(slide, filename, x, y, w, h):
        path = REPORTS / "figures" / filename
        with Image.open(path) as im:
            ratio = im.width / im.height
        dw, dh = min(w, h * ratio), min(h, w / ratio)
        slide.shapes.add_picture(
            str(path),
            Inches(x + (w - dw) / 2),
            Inches(y + (h - dh) / 2),
            width=Inches(dw),
            height=Inches(dh),
        )

    def new(title, source, notes):
        slide = deck.slides.add_slide(layout)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(WHITE)
        box(slide, 0, 0, 13.333, 0.12, TEAL)
        text(slide, "IE313 • NHÓM 10 • AIRBNB BANGKOK", 0.55, 0.30, 11, 0.3, 11, MUTED)
        text(slide, title, 0.55, 0.84, 12.2, 0.8, 29, NAVY, True)
        box(slide, 0.55, 6.88, 12.2, 0.012, "CFDBE1")
        text(slide, source, 0.55, 7.00, 11.6, 0.32, 9, MUTED)
        text(slide, f"{len(deck.slides):02}", 12.15, 7.00, 0.6, 0.3, 10, TEAL, True)
        slide.notes_slide.notes_text_frame.text = notes
        return slide

    def card(slide, x, y, w, title, body):
        box(slide, x, y, w, 1.8, PALE)
        text(slide, title, x + 0.18, y + 0.13, w - 0.36, 0.42, 23, TEAL, True)
        text(slide, body, x + 0.18, y + 0.70, w - 0.36, 1.02, 17)

    s = new(
        "Phân tích và trực quan hóa giá thuê Airbnb",
        "Inside Airbnb • Bangkok 29/06/2026 • CC BY 4.0 • Bản trình bày 12/10/2026",
        "",
    )
    text(s, "Bangkok và các yếu tố liên quan", 0.58, 1.62, 12, 0.8, 31, TEAL, True)
    card(
        s,
        0.58,
        2.76,
        5.95,
        "Phạm Quốc Thắng • 25410304",
        "RQ, EDA, biểu đồ, diễn giải\nBáo cáo và trình bày",
    )
    card(
        s,
        6.8,
        2.76,
        5.95,
        "Lê Minh Thiện • 25410313",
        "Kiểm toán, làm sạch, thống kê\nTrain/CV và kiểm tra tái lập",
    )
    text(
        s,
        "Minh bạch: dữ liệu Inside Airbnb; không tham khảo dự án/code mẫu; không dùng đề tài/dữ liệu cho môn khác. Có AI hỗ trợ triển khai và rà soát; nhóm chịu trách nhiệm kiểm tra.",
        0.63,
        5.03,
        12.1,
        1.4,
        20,
    )

    s = new(
        "Bốn câu hỏi dẫn dắt phân tích",
        "Đối chiếu: SPECIFICATION.md • Không suy ra quan hệ nhân quả",
        "",
    )
    card(
        s,
        0.6,
        1.95,
        5.95,
        "01  Phân bố giá",
        "Mức giá điển hình và độ phân tán?\nHistogram, median, IQR",
    )
    card(
        s,
        6.8,
        1.95,
        5.95,
        "02  Loại phòng",
        "Giá khác nhau giữa các loại phòng?\nBoxplot, thống kê nhóm, ANOVA thăm dò",
    )
    card(
        s,
        0.6,
        4.12,
        5.95,
        "03  Quy mô chỗ ở",
        "Số khách/phòng/giường liên hệ với giá?\nPearson r, p, n và scatter",
    )
    card(
        s,
        6.8,
        4.12,
        5.95,
        "04  Khu vực",
        "Khác biệt trong cùng loại phòng?\nPivot median/count/IQR, bar, heatmap",
    )

    s = new(
        "Dữ liệu: một listing trong một đợt thu thập",
        "Nguồn: data/source_manifest.json • reports/audit/bangkok_columns.csv",
        "",
    )
    card(s, 0.6, 1.95, 3.95, "31.069 × 90", "Dòng × cột trong raw\n31.069 ID duy nhất")
    card(s, 4.7, 1.95, 3.95, "28.987 × 16", "Baseline sau làm sạch\n12 biến chính + 4 cột truy vết")
    card(s, 8.8, 1.95, 3.95, "THB / đêm", "Báo giá niêm yết\nChưa xác minh đủ thuế/phí")
    text(
        s,
        "Biến chính: price; room_type, property_type, neighbourhood_cleansed, host_is_superhost; accommodates, bedrooms, bathrooms, beds, minimum_nights, number_of_reviews, review_scores_rating.",
        0.7,
        4.3,
        11.8,
        1.4,
        21,
    )
    text(
        s,
        "7.231 host • 3.215 host có nhiều listing → quan sát có thể phụ thuộc.",
        0.7,
        6.0,
        12,
        0.5,
        21,
        TEAL,
        True,
    )

    s = new(
        "Làm sạch có nhật ký, giữ nguyên dữ liệu thô",
        "Nguồn: docs/CLEANING_PLAN.md • reports/cleaning/ • notebook 01",
        "",
    )
    picture(s, "report_analysis_flow.png", 0.6, 1.9, 12.1, 1.55)
    text(
        s,
        "2.082 dòng thiếu giá: loại bỏ, không điền giá mục tiêu.\n2.730 ngoại lệ IQR: đánh dấu để rà soát, giữ trong phân tích.\nMissing predictor: giữ ở baseline, công bố n cho mỗi phép tính.\nCác cột báo giá/revenue, ID và host_id: không làm predictor.",
        0.8,
        3.85,
        11.8,
        2.65,
        23,
    )

    s = new(
        "RQ1 • Trung bình cao hơn nhiều so với trung vị",
        "Nguồn: reports/stats/rq1_price_distribution.json • fig1_price_distribution.png",
        "",
    )
    picture(s, "fig1_price_distribution.png", 0.45, 1.93, 9.35, 4.5)
    text(s, "1.592", 10.0, 2.2, 2.8, 0.8, 35, TEAL, True)
    text(s, "THB/đêm\nTrung vị", 10, 3.0, 2.75, 1, 20)
    text(
        s,
        "2.955,96\nTrung bình\n\n777 điểm ngoài khung 10.000 vẫn được giữ.",
        10,
        4.1,
        2.75,
        2.5,
        20,
    )

    s = new(
        "RQ2 • So sánh giá cần xét loại phòng",
        "Nguồn: rq2_room_type_stats.json • rq2_anova_condition_checks.json • notebook 02–03",
        "",
    )
    picture(s, "fig2_price_by_room_type.png", 0.45, 1.9, 9.35, 4.42)
    text(s, "431,22", 10, 2.05, 2.75, 0.7, 31, TEAL, True)
    text(s, "Shared room\nMedian THB/đêm", 10, 2.85, 2.75, 1.1, 18)
    text(s, "1.715,44", 10, 4.1, 2.75, 0.7, 31, TEAL, True)
    text(s, "Entire home/apt\nMedian THB/đêm", 10, 4.9, 2.75, 1.1, 18)
    text(
        s,
        "ANOVA thăm dò: F=5,9323; p=0,000485542. Không suy ra tác động nhân quả.",
        0.8,
        6.33,
        12,
        0.45,
        17,
    )

    s = new(
        "RQ3 • p nhỏ, nhưng tương quan tuyến tính yếu",
        "Nguồn: reports/stats/rq3_size_correlation.json • fig3_price_vs_size_scatter.png",
        "",
    )
    picture(s, "fig3_price_vs_size_scatter.png", 0.55, 1.75, 7.7, 5.0)
    columns = [
        ("Biến\nSố khách\nPhòng ngủ\nPhòng tắm\nSố giường", 8.55, 1.65),
        ("r\n0,1067\n0,0881\n0,0841\n0,0688", 10.3, 1.1),
        ("n\n28.987\n25.428\n26.286\n27.717", 11.55, 1.18),
    ]
    for value, x, width in columns:
        text(s, value, x, 2.15, width, 2.85, 19)
    text(
        s,
        "p lần lượt: 4,268×10⁻⁷⁴; 6,036×10⁻⁴⁵; 1,977×10⁻⁴²; 1,902×10⁻³⁰.",
        8.6,
        5.15,
        4.15,
        1.0,
        16,
        MUTED,
    )
    text(s, "Không dùng EDA toàn mẫu để chọn biến mô hình.", 8.6, 6.16, 4.15, 0.57, 17, TEAL, True)

    s = new(
        "RQ4 • Khác biệt khu vực trong cùng loại phòng",
        "Nguồn: reports/stats/rq4_neighbourhood_* • fig4_price_top_neighbourhoods_bar.png",
        "",
    )
    picture(s, "fig4_price_top_neighbourhoods_bar.png", 0.45, 1.8, 9.3, 4.9)
    text(s, "678,03 → 3.115,50", 9.9, 2.05, 2.9, 1.15, 25, TEAL, True)
    text(
        s,
        "THB/đêm\nBang Phlat → Samphanthawong\nCùng Entire home/apt\nn≥20: chỉ để chọn nhóm hiển thị.",
        9.9,
        3.42,
        2.9,
        3.1,
        19,
    )

    s = new(
        "Hồi quy: chưa có biến qua chính sách lựa chọn",
        "Nguồn: reports/modeling/feature_selection.json • model_comparison.csv • notebook 04",
        "",
    )
    card(s, 0.6, 1.95, 3.95, "20.290 / 8.697", "Train / test • seed 42\nCV 4-fold trên train")
    card(s, 4.7, 1.95, 3.95, "0 biến số đạt", "|r|≥0,30 và p<0,05\nKiểm tra từng training fold")
    card(
        s,
        8.8,
        1.95,
        3.95,
        "ANOVA làm bộ lọc",
        "Chưa xác nhận điều kiện\nLưu chẩn đoán thực từng fold",
    )
    text(
        s,
        "Mốc trung bình train: 2.976,12 THB/đêm\nTest MSE = 315.021.203,83 THB² • R² = −0,000014334",
        0.75,
        4.22,
        12,
        1.25,
        25,
        TEAL,
        True,
    )
    text(
        s,
        "Đây là giới hạn của chính sách chọn biến hiện tại.\n82,81% listing test thuộc host đã xuất hiện ở train; chưa đánh giá host mới.",
        0.75,
        5.83,
        12,
        0.94,
        20,
    )

    s = new(
        "Kết quả và giới hạn cần giữ khi diễn giải",
        "Nguồn: báo cáo Word, mục Kết quả phân tích / Kết luận • Phần trình bày: 10 phút",
        "",
    )
    text(
        s,
        "Giá lệch phải → đọc median cùng IQR.\nLoại phòng khác nhau → so sánh đúng phân khúc.\nQuy mô–giá có r nhỏ → p nhỏ không thay thế độ lớn liên hệ.\nKhu vực khác nhau → đối chiếu trong cùng loại phòng.",
        0.72,
        1.98,
        12,
        3.0,
        25,
    )
    box(s, 0.6, 5.35, 12.1, 1.2, PALE)
    text(
        s,
        "Giới hạn: một snapshot, missing predictor, cực trị chưa xác minh, phụ thuộc theo host và thuế/phí chưa đầy đủ. Không suy ra nhân quả.",
        0.82,
        5.55,
        11.65,
        0.92,
        21,
    )

    s = new(
        "Demo 5 phút • Từ nguồn đến kết quả",
        "Hướng dẫn đầy đủ: reports/Nhom10_Kich_ban_thuyet_trinh.md • docs/RUNBOOK.md",
        "",
    )
    text(
        s,
        "00:00  Xác minh nguồn và checksum raw\n00:45  Notebook 01: 31.069 → 28.987 listing\n01:30  Notebook 02–03: r/p/n và hình từ dữ liệu thực\n02:30  Notebook 04: train/CV, chọn biến và benchmark\n03:45  Đối chiếu một bảng, một hình trong báo cáo\n04:30  Giải thích giới hạn và trả lời câu hỏi",
        0.85,
        1.98,
        11.7,
        4.6,
        25,
    )

    s = new(
        "Trao đổi và câu hỏi",
        "Inside Airbnb: insideairbnb.com/get-the-data/ • Nguồn phương pháp đầy đủ trong báo cáo và docs/",
        "",
    )
    text(
        s,
        "Vì sao giữ giá cực trị?\nVì sao p rất nhỏ nhưng r vẫn yếu?\nVì sao chưa xây hồi quy Linear/Ridge?\nTrain/test ngẫu nhiên có đánh giá host mới không?",
        0.85,
        2.05,
        11.7,
        3.5,
        28,
    )
    text(s, "Nhóm 10 • Phạm Quốc Thắng & Lê Minh Thiện", 0.85, 6.1, 11.7, 0.5, 22, TEAL, True)
    notes = [
        "Thắng – 45 giây. Giới thiệu đề tài và hai thành viên. Phân tích báo giá niêm yết tại một snapshot, không phải giá giao dịch. Nhóm xác nhận không tham khảo code/dự án mẫu, không sử dụng đề tài/dữ liệu cho môn khác; có AI hỗ trợ và nhóm chịu trách nhiệm kiểm tra.",
        "Thắng – 40 giây. Nêu bốn RQ. Luồng trình bày: dữ liệu, xử lý, bốn RQ, kiểm tra khả thi hồi quy, kết quả và giới hạn.",
        "Thiện – 50 giây. last_scraped từ 29/06 đến 02/07/2026. Một dòng là listing, không phải listing–ngày. Giá THB/đêm được đối chiếu trên 28.987 cặp với trường báo giá; chưa xác minh toàn bộ thuế/phí. Không suy ra doanh thu.",
        "Thiện – 50 giây. Loại 2.082 dòng thiếu giá, không suy diễn y. Ngoại lệ IQR chỉ được đánh dấu; giữ giá dương hữu hạn. Baseline giữ missing predictor; EDA dùng cặp hợp lệ. Imputer/scaler chỉ fit train nếu có mô hình khả thi. Raw bất biến, kiểm tra SHA-256.",
        "Thắng – 65 giây. Median 1.592; mean 2.955,96 THB/đêm; Q1 1.041,465 và Q3 2.539,565. Phân phối lệch phải. Trục log chỉ để hiển thị, không biến đổi target. Panel phải có 28.210 dòng trong khung; 777 dòng ngoài khung vẫn được giữ trong thống kê.",
        "Thắng – 70 giây. Shared room median 431,22 n=498; Private room 1.389,83 n=8.914; Entire home/apt 1.715,44 n=19.300; Hotel room 1.999 n=275. Hộp/râu tính toàn nhóm trước khi giới hạn trục 8.000. Q1–Q3 không là khoảng tin cậy. ANOVA F=5,9323 p=0,000485542; Levene W=3,123 p=0,0247826. Kết quả thăm dò, có giới hạn phân phối lệch, phương sai và phụ thuộc theo host; không suy ra nhân quả.",
        "Thắng – 65 giây. r chỉ khoảng 0,069–0,107; p nhỏ với n lớn không đồng nghĩa dự báo tốt. Scatter lấy tối đa 3.000 điểm mỗi biến và giới hạn trục 20.000; OLS/Pearson tính toàn bộ cặp hợp lệ. Không dùng EDA toàn mẫu để chọn biến trong CV. Chưa có phân tích độ nhạy để quy r thấp cho cực trị.",
        "Thắng – 60 giây. Trong cùng Entire home/apt: Bang Phlat có n=284, median 678,03; Samphanthawong n=105, median 3.115,50 THB/đêm. Ngưỡng n≥20 chỉ phục vụ hiển thị. Pivot 50×4 có 43 ô không mẫu và 27 ô với 1–4 listing. Chưa kiểm soát chất lượng/tiện nghi nên không tách riêng tác động vị trí.",
        "Thiện – 70 giây. Split 70/30 seed42, KFold 4-fold shuffle42. Chọn biến/fit transform chỉ trên training fold. Không biến số nào đạt |r|≥0,30 và p<0,05; điều kiện ANOVA làm bộ lọc phân loại chưa xác nhận. Chưa xây hồi quy theo chính sách hiện tại, không chứng minh dữ liệu không thể mô hình hóa. Dummy dự đoán mean train 2.976,12; test MSE 315.021.203,83 và R² −0,000014334. Không có hệ số, alpha tối ưu hoặc hình phần dư hồi quy thực. 82,81% listing test thuộc host đã thấy ở train.",
        "Thắng – 85 giây. Tổng hợp bốn phát hiện và các giới hạn, đối chiếu mỗi phát hiện với hình tương ứng. Nêu các việc chưa có cơ sở: đề xuất giá tối ưu, doanh thu, nhân quả, dự báo thời điểm khác hoặc chủ nhà mới. Kết quả hỗ trợ chọn nhóm báo giá để so sánh. Tổng slide 1–10: 600 giây. Chuyển Thiện thực hiện demo.",
        "Thiện chủ trì demo, Thắng mở báo cáo. 0:00–0:45 manifest/hash raw; 0:45–1:30 notebook 01 và cleaning; 1:30–2:30 notebook 02–03 với r/p/n và hình; 2:30–3:45 chạy notebook 04 từ kernel mới và mở artifact; 3:45–4:30 đối chiếu Word/PDF; 4:30–5:00 giới hạn. Đây là kịch bản dự kiến; cần tổng duyệt trên máy trình bày.",
        "Hai thành viên cùng trả lời. Không nói dữ liệu không thể mô hình hóa, r thấp do outlier hay vị trí gây giá cao khi chưa có kiểm chứng. Chỉ nguồn tại manifest, EDA_METHODS, MODELING_METHODS, bốn notebook và review Bước 7. Câu hỏi ngoài phạm vi được ghi nhận là hạn chế hoặc hướng nghiên cứu tiếp theo.",
    ]
    for slide, note in zip(deck.slides, notes, strict=True):
        slide.notes_slide.notes_text_frame.text = note
    path = REPORTS / "Nhom10_Slides.pptx"
    deck.save(path)
    print(path)


if __name__ == "__main__":
    main()
