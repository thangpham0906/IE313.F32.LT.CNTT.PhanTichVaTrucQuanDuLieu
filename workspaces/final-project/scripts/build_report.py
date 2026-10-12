"""Hoàn thiện báo cáo từ bản khung template và artifact đã review của Nhóm 10."""

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
REPORTS = ROOT / "reports"


def load(name):
    """Đọc kết quả thực nghiệm đã sinh từ notebook."""
    return json.loads((REPORTS / name).read_text())


def vn(value, digits=2):
    """Định dạng số theo dấu phân cách tiếng Việt chỉ ở lớp trình bày."""
    return f"{value:,.{digits}f}".replace(",", "_").replace(".", ",").replace("_", ".")


def main():
    """Giữ trang bìa/style gốc; tạo thân bài 8 trang theo ngắt trang có chủ đích."""
    doc = Document(REPORTS / "archive/Nhom10_Bao_cao_Khung_20261010.docx")
    first = next(p for p in doc.paragraphs if p.style.name == "Heading 1")
    body = doc._element.body
    for element in list(body)[list(body).index(first._p) :]:
        if element.tag != qn("w:sectPr"):
            body.remove(element)
    for p in doc.paragraphs:
        if "BẢN KHUNG" in p.text:
            p.text = "BÁO CÁO ĐỒ ÁN MÔN HỌC\nKết quả thực nghiệm cập nhật ngày 12/10/2026"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(12)
                run.font.color.rgb = RGBColor.from_string("333333")
                run.font.highlight_color = None
    for section in doc.sections:
        for part in (section.header, section.first_page_header):
            for paragraph in part.paragraphs:
                for run in paragraph.runs:
                    run.text = run.text.replace(" | BẢN KHUNG", "")
    for name in ["Heading 1", "Heading 2"]:
        doc.styles[name].paragraph_format.keep_with_next = True
        doc.styles[name].font.color.rgb = RGBColor.from_string("000000")
    price = load("stats/rq1_price_distribution.json")
    anova = load("stats/rq2_anova_condition_checks.json")
    corr = load("stats/rq3_size_correlation.json")
    model = load("modeling/modeling_metadata.json")
    selection = load("modeling/feature_selection.json")
    bench = model["benchmark_results"]
    overlap = model["split_info"]["host_overlap"]
    with (ROOT / "data/processed/bangkok_cleaned.csv").open() as handle:
        records = list(csv.DictReader(handle))

    def p(text, style="BT"):
        par = doc.add_paragraph(text, style)
        par.paragraph_format.widow_control = True
        return par

    def heading(text, level=1):
        return p(text, f"Heading {level}")

    def page():
        doc.add_page_break()

    def table(title, headers, rows, widths=None):
        caption = p(title, "Caption")
        caption.paragraph_format.keep_with_next = True
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.autofit = False
        if widths:
            for column, width in zip(t.columns, widths, strict=True):
                column.width = Cm(width)
        for i, value in enumerate(headers):
            t.rows[0].cells[i].text = str(value)
        for values in rows:
            for cell, value in zip(t.add_row().cells, values, strict=True):
                cell.text = str(value)
        repeat = OxmlElement("w:tblHeader")
        t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row_no, row in enumerate(t.rows):
            trpr = row._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            for i, cell in enumerate(row.cells):
                if widths:
                    cell.width = Cm(widths[i])
                if row_no == 0:
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:fill"), "E9F0F5")
                    cell._tc.get_or_add_tcPr().append(shd)
                for par in cell.paragraphs:
                    par.style = doc.styles["BT"]
                    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    par.paragraph_format.first_line_indent = Cm(0)
                    par.paragraph_format.left_indent = Cm(0)
                    par.paragraph_format.right_indent = Cm(0)
                    par.paragraph_format.space_before = Pt(1)
                    par.paragraph_format.space_after = Pt(2)
                    par.paragraph_format.line_spacing = 1.0
                    for run in par.runs:
                        run.font.size = Pt(10.5)
                        run.bold = row_no == 0
        return t

    def figure(filename, caption, width=16):
        par = p("")
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(0)
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.keep_with_next = True
        par.add_run().add_picture(str(REPORTS / "figures" / filename), width=Cm(width))
        cap = p(caption, "Caption")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(4)
        for run in cap.runs:
            run.font.size = Pt(10)

    # Trang thân 1: đúng hai đoạn Giới thiệu.
    heading("GIỚI THIỆU")
    p(
        f"Đề tài phân tích và trực quan hóa giá thuê Airbnb tại Bangkok, tập trung vào phân bố giá và mối liên hệ với loại phòng, quy mô chỗ ở, khu vực. Nhóm sử dụng Python, Pandas, SciPy và Matplotlib để kiểm toán, làm sạch, thống kê và trả lời bốn câu hỏi nghiên cứu; đồng thời đánh giá khả năng chọn biến cho hồi quy trên train/CV. Từ 31.069 listing, giữ 28.987 quan sát có giá; trung vị là {vn(price['median'])} THB/đêm. Giá lệch phải, trung vị khác nhau theo loại phòng và khu vực; tương quan tuyến tính quy mô–giá có độ lớn nhỏ."
    )
    p(
        "Bộ dữ liệu được tham khảo từ Inside Airbnb, bản Bangkok ngày 29/06/2026, tải ngày 10/10/2026 theo giấy phép CC BY 4.0 [1, 2]. Nhóm xây dựng quy trình riêng cho đồ án IE313, có hỗ trợ công cụ AI trong triển khai, soạn thảo và rà soát; nhóm chịu trách nhiệm kiểm tra kết quả. Nhóm xác nhận không tham khảo dự án/code mẫu, chưa từng và hiện không đồng thời sử dụng đề tài/bộ dữ liệu này làm đồ án môn khác. Nguồn dữ liệu, thư viện và phương pháp tham khảo được trích dẫn trong báo cáo."
    )
    heading("MÔ TẢ BỘ DỮ LIỆU")
    heading("Nguồn và đơn vị quan sát", 2)
    p(
        "Mỗi dòng là một listing duy nhất trong một đợt dữ liệu Bangkok, không phải giao dịch thuê hay chuỗi listing–ngày. Tệp gốc có 31.069 dòng, 90 cột và 31.069 ID duy nhất; ngày last_scraped từ 29/06 đến 02/07/2026. Giá được xác minh là báo giá theo đêm bằng THB qua từ điển nguồn và 28.987 cặp khớp với price_quote_price_per_night [1, 2]. Chưa xác minh đầy đủ thuế/phí cuối cùng."
    )
    table(
        "Bảng 1. Quy mô và chất lượng dữ liệu",
        ["Nội dung", "Kết quả"],
        [
            ("Nguồn dữ liệu", "Inside Airbnb – Bangkok, bản 29/06/2026"),
            ("Tệp gốc", "31.069 dòng × 90 cột; 20.625.983 bytes"),
            ("Thiếu giá / trùng ID", "2.082 dòng thiếu giá (6,70%); không trùng ID"),
            ("Dữ liệu phân tích", "28.987 dòng × 16 cột; 12 biến chính và 4 cột truy vết"),
            (
                "Phạm vi suy rộng",
                "Báo giá niêm yết tại một snapshot; không phải doanh thu hay giá giao dịch",
            ),
        ],
        [4.2, 11.8],
    )

    page()
    heading("Các biến sử dụng trong phân tích", 2)
    p(
        "Trong 12 biến chính có 8 biến số (kể cả price) và 4 biến phân loại. ID, host_id, last_scraped và price_raw được giữ để truy vết; không dùng làm predictor. Bảng 2 ghi phạm vi quan sát sau làm sạch, không phải giới hạn lý thuyết. Danh mục toàn bộ 90 cột và nguồn giải thích nằm trong bangkok_columns.csv nộp kèm."
    )
    definitions = [
        ("price", "float64", "Báo giá THB/đêm; biến mục tiêu"),
        ("room_type", "Phân loại", "Loại phòng theo nguồn"),
        ("property_type", "Phân loại", "Loại hình chỗ ở theo nguồn"),
        ("neighbourhood_cleansed", "Phân loại", "Khu vực/quận chuẩn hóa của nguồn"),
        ("accommodates", "float64", "Số khách tối đa"),
        ("bedrooms", "float64", "Số phòng ngủ; 0 có thể là studio"),
        ("bathrooms", "float64", "Số phòng tắm; cho phép phần lẻ"),
        ("beds", "float64", "Số giường"),
        ("minimum_nights", "float64", "Số đêm tối thiểu"),
        ("number_of_reviews", "float64", "Số lượt đánh giá"),
        ("review_scores_rating", "float64", "Điểm đánh giá tổng thể"),
        ("host_is_superhost", "Phân loại", "Trạng thái superhost: t/f"),
    ]
    rows = []
    for i, (name, dtype, meaning) in enumerate(definitions, 1):
        values = [r[name] for r in records if r[name] != ""]
        extent = (
            f"{len(set(values))} nhóm"
            if dtype == "Phân loại"
            else f"{vn(min(map(float, values)), 2)}–{vn(max(map(float, values)), 2)}"
        )
        rows.append((i, name, dtype, extent, meaning))
    table(
        "Bảng 2. Mô tả 12 biến chính trong baseline",
        ["STT", "Tên cột", "Kiểu dữ liệu", "Phạm vi", "Giải thích"],
        rows,
        [1.1, 4.1, 2.1, 3.5, 5.2],
    )
    p(
        "Giá có từ 1 đến 1.320.255 THB/đêm; các mức cực đoan được giữ và đánh dấu để rà soát, chưa quy kết là lỗi hoặc căn cao cấp. Predictor còn thiếu: bedrooms 3.559 dòng, bathrooms 2.701, beds 1.270, minimum_nights 2 và review_scores_rating 9.451. Phân tích từng cặp công bố n hợp lệ; baseline không điền khuyết hàng loạt."
    )
    p(
        "Dữ liệu sạch có 7.231 host, trong đó 3.215 host có nhiều listing. Vì các quan sát có thể cùng chủ nhà, các p-value trong báo cáo được diễn giải thăm dò. Không suy ra quan hệ nhân quả từ tương quan hoặc chênh lệch nhóm."
    )

    page()
    heading("PHƯƠNG PHÁP PHÂN TÍCH")
    heading("Quy trình và xử lý dữ liệu", 2)
    p(
        "Quy trình thực tế ở Hình 1 vận dụng Bài 01–07 [3]. Bốn notebook tách kiểm toán/làm sạch, EDA, trực quan và đánh giá khả thi mô hình. Dữ liệu thô giữ nguyên checksum; mọi bảng, hình và nhận xét đều truy về kết quả chạy."
    )
    figure(
        "report_analysis_flow.png",
        "Hình 1. Quy trình nhóm triển khai; nhánh mô hình được đánh giá khả thi trên train/CV.",
    )
    table(
        "Bảng 3. Giải trình các quyết định làm sạch",
        ["Vấn đề", "Xử lý", "Lý do và ảnh hưởng"],
        [
            (
                "Giá dạng chuỗi",
                "Parse tiền tệ, giữ giá dương hữu hạn",
                "Loại 2.082 dòng thiếu giá; không suy diễn y",
            ),
            (
                "Ngoại lệ IQR",
                "Đánh dấu, giữ nguyên",
                "2.730 dòng được lưu để rà soát; không cắt theo mức giá tùy ý",
            ),
            (
                "Predictor khuyết/sai",
                "Chuẩn hóa kiểu; giữ khuyết",
                "EDA dùng cặp hợp lệ; imputer chỉ fit trong train nếu mô hình khả thi",
            ),
            (
                "Định danh và cột báo giá",
                "Giữ cột truy vết; chặn cột lộ y",
                "Không có ID trùng; price_quote/revenue không làm predictor",
            ),
        ],
        [3.1, 4.0, 8.9],
    )
    heading("Thống kê và trực quan hóa", 2)
    p(
        "RQ1 dùng mean/median, Q1–Q3, độ lệch chuẩn mẫu (ddof=1), IQR và hình dạng phân phối. Skewness/Fisher excess kurtosis EDA dùng bias=False. RQ2 mô tả nhóm và ANOVA thăm dò với F, p. RQ3 dùng Pearson r, p và n theo cặp. RQ4 dùng pivot median/count/IQR theo quận × loại phòng, rồi so sánh trong cùng loại phòng."
    )
    p(
        "Biểu đồ dùng Matplotlib OO API, nhãn tiếng Việt và ảnh 300 DPI. Thang log trên trục giá chỉ phục vụ đọc histogram toàn miền; không biến đổi giá trong thống kê/mô hình. Giới hạn trục ở boxplot/scatter không thay tập mẫu dùng tính phân vị, Pearson hoặc đường OLS."
    )
    p(
        "[MỞ RỘNG NGOÀI BÀI HỌC] Levene center=median bổ sung chẩn đoán phương sai [4]; p≥0,05 không chứng minh phương sai bằng nhau. pd.qcut chỉ minh họa phân khoảng theo phân vị trên bản sao EDA, bên cạnh pd.cut của Bài 02 [5]; không tạo predictor từ price. Hai kỹ thuật có nguồn và giới hạn trong EDA_METHODS.md."
    )

    page()
    heading("PHÂN TÍCH THĂM DÒ VÀ TRỰC QUAN HÓA")
    heading("RQ1. Giá phân bố như thế nào?", 2)
    figure(
        "fig1_price_distribution.png",
        "Hình 2. Histogram toàn miền và vùng ≤10.000 THB; trung bình/trung vị tính toàn mẫu.",
    )
    p(
        f"Giá trung vị {vn(price['median'])} THB/đêm, trung bình {vn(price['mean'])}; Q1–Q3 là {vn(price['q1'])}–{vn(price['q3'])}. Skewness {vn(price['skewness'], 4)}, excess kurtosis {vn(price['kurtosis'], 4)} cho thấy phân phối lệch phải mạnh. Có 2.730 ngoại lệ IQR (9,42%); 777 điểm ngoài vùng giá 10.000 là tiêu chí hiển thị khác với ngoại lệ IQR."
    )
    heading("RQ2. Giá khác nhau theo loại phòng?", 2)
    figure(
        "fig2_price_by_room_type.png",
        "Hình 3. Boxplot và trung vị/Q1–Q3 tính toàn nhóm; trục boxplot giới hạn 8.000 THB.",
    )
    p(
        "Trung vị Entire home/apt 1.715,44 (n=19.300), Hotel room 1.999,00 (n=275), Private room 1.389,83 (n=8.914), Shared room 431,22 THB/đêm (n=498). Thanh Q1–Q3 mô tả phân tán, không phải khoảng tin cậy. Quy mô nhóm khác nhau đáng kể."
    )
    omnibus = anova["omnibus_test"]
    lev = anova["condition_checks"]["homoscedasticity_evaluation"]
    p(
        f"ANOVA thăm dò F={vn(omnibus['f_statistic'], 4)}, p={omnibus['p_value']:.6g}. Levene W={vn(lev['statistic'], 4)}, p={lev['p_value']:.6g}. Phân phối lệch và host clustering hạn chế suy luận; không tiếp tục so sánh cặp nhóm hoặc kết luận tác động nhân quả."
    )

    page()
    heading("RQ3. Quy mô chỗ ở liên hệ thế nào với giá?", 2)
    p(
        "Bốn biến quy mô gồm số khách, phòng ngủ, phòng tắm và giường. Mỗi cặp chỉ dùng quan sát hữu hạn hợp lệ; các số r/p này thuộc EDA toàn mẫu, không dùng trực tiếp để chọn biến cho mô hình."
    )
    figure(
        "fig3_price_vs_size_scatter.png",
        "Hình 4. Scatter quy mô–giá; lấy mẫu để vẽ, Pearson và OLS tính toàn bộ cặp hợp lệ.",
        15.5,
    )
    table(
        "Bảng 4. Pearson giữa quy mô và giá, kết quả thăm dò",
        ["Biến", "n hợp lệ", "r", "p-value"],
        [
            (r["feature"], vn(r["n_valid"], 0), vn(r["pearson_r"], 4), f"{r['p_value']:.4g}")
            for r in corr
        ],
        [5, 3, 3, 5],
    )
    p(
        "Các tương quan dương có độ lớn nhỏ (r≈0,069–0,107). p nhỏ không đồng nghĩa liên hệ mạnh hoặc hữu ích để dự báo. Chưa thực hiện phân tích độ nhạy nên không quy nguyên nhân r thấp cho giá cực trị. Missing theo cặp và sự phụ thuộc theo host tiếp tục là giới hạn."
    )
    p(
        "Mỗi panel chọn tối đa 3.000 điểm với seed42 và chỉ hiển thị giá ≤20.000 THB; số điểm thực vẽ lần lượt 2.974, 2.962, 2.978 và 2.975. Đường OLS vẫn fit toàn cặp hợp lệ, không fit lại trên phần giá đang hiển thị. Số liệu chi tiết nằm trong figures_catalog.json."
    )

    page()
    heading("RQ4. Giá khác nhau theo khu vực?", 2)
    p(
        "Bảng chéo gồm 50 quận × 4 loại phòng, có 43 ô không có mẫu và 27 ô có 1–4 listing. Nhóm đối chiếu trung vị trong cùng loại phòng để tránh gộp các cơ cấu phòng khác nhau; tên quận giữ theo nguồn, không suy diễn thành khu du lịch hoặc trung tâm khi chưa kiểm chứng."
    )
    figure(
        "fig4_price_top_neighbourhoods_bar.png",
        "Hình 5. Trung vị theo quận trong cùng Entire home/apt; chỉ hiển thị nhóm n≥20.",
        15.5,
    )
    table(
        "Bảng 5. Một số khu vực trong cùng loại Entire home/apt",
        ["Quận", "n", "Trung vị THB/đêm"],
        [
            ("Bang Phlat", 284, "678,03"),
            ("Thon buri", 396, "940,14"),
            ("Bang Na", 586, "969,48"),
            ("Phra Nakhon", 130, "2.918,50"),
            ("Lat Phrao", 31, "3.010,00"),
            ("Samphanthawong", 105, "3.115,50"),
        ],
        [7, 3, 6],
    )
    p(
        "Ngưỡng n≥20 là lựa chọn hiển thị, không bảo đảm tính đại diện hay loại bỏ sai số mẫu. Chênh lệch trung vị mô tả mẫu niêm yết đang có; chưa kiểm soát đồng thời chất lượng chỗ ở, tiện nghi và các yếu tố khác nên không coi đây là tác động riêng của vị trí."
    )
    p(
        "Heatmap bổ sung trong thư mục hình cho 15 quận đông listing hiển thị 23.043 quan sát, ghi trung vị và n từng ô; ô ít mẫu có dấu *. Bảng median/count/IQR đầy đủ được nộp kèm để đối chiếu các quận ngoài Hình 5."
    )

    page()
    heading("KẾT QUẢ PHÂN TÍCH")
    heading("Hồi quy bổ trợ: đánh giá khả thi theo Bài 06–07", 2)
    p(
        f"Chia 20.290 listing train và 8.697 test, seed42; KFold 4-fold shuffle42 trên train. Chọn biến, imputer/scaler/encoder chỉ fit trên tập train tương ứng. {overlap['overlapping_hosts_count']:,} host xuất hiện ở cả hai tập; {overlap['test_listings_belonging_to_train_hosts']:,} listing test (82,81%) thuộc host đã thấy. Thiết kế này không đánh giá khả năng dự báo cho chủ nhà mới. [3, 6]".replace(
            "2,229", "2.229"
        ).replace("7,202", "7.202")
    )
    p(
        "Ngưỡng chọn biến số của repository là |r|≥0,30 và p<0,05, sử dụng trị chưa làm tròn trong mỗi training fold. Cả 7 biến số đều không qua ngưỡng ở 4 fold và toàn train. Với phân loại, nhóm chưa có cơ sở dùng ANOVA làm bộ lọc do giả định suy luận chưa xác nhận; chẩn đoán thực của từng nhóm/fold được lưu lại, không gán sẵn mọi giả định đều vi phạm."
    )
    table(
        "Bảng 6. Kiểm tra chọn biến trên toàn train",
        ["Biến số", "n hợp lệ", "Pearson r", "p-value", "Qua tiêu chí"],
        [
            (
                name,
                vn(row["n_valid"], 0),
                vn(row["pearson_r"], 5),
                f"{row['p_value']:.4g}",
                "Không" if not row["selected"] else "Có",
            )
            for name, row in selection["full_train"]["numeric_features"].items()
        ],
        [4.7, 2.3, 2.6, 3, 3.4],
    )
    p(
        "Chưa xây hồi quy Linear/Ridge theo chính sách chọn biến hiện tại; kết quả không chứng minh dữ liệu không thể mô hình hóa. Không có hệ số, alpha tối ưu hay chẩn đoán hồi quy cho Bangkok. Nhánh phần mềm khả thi đã kiểm thử trên dữ liệu tổng hợp, không dùng làm kết quả đồ án."
    )
    heading("Mốc dự đoán trung bình", 2)
    p(
        f"DummyRegressor dự đoán cố định trung bình train {vn(bench['train_mean_prediction_thb'])} THB/đêm; mỗi fold có trung bình riêng. Đây là API bổ sung kỹ thuật theo repository, liên hệ công thức R² trong Bài 06; không thay thế hồi quy có predictor. R² = 1 − SSE/SST, với SST dùng trung bình của tập đang đánh giá [6]."
    )
    table(
        "Bảng 7. Kết quả mốc trung bình",
        ["Tập đánh giá", "MSE (THB²)", "R²"],
        [
            ("Train", vn(bench["train_mse"]), f"{bench['train_r2']:.7f}"),
            (
                "CV mean ± std",
                f"{vn(bench['cv_4fold_mse_mean'])} ± {vn(bench['cv_4fold_mse_std'])}",
                f"{bench['cv_4fold_r2_mean']:.7f} ± {bench['cv_4fold_r2_std']:.7f}",
            ),
            ("Test", vn(bench["test_mse"]), f"{bench['test_r2']:.9f}"),
        ],
        [3.3, 8.2, 4.5],
    )
    p(
        "CV std dùng ddof=0. R² test hơi âm: dự đoán trung bình train có sai số nhỉnh hơn trung bình test trong SST. MSE giữa các fold dao động mạnh; chưa đánh giá dự báo cho host mới."
    )

    page()
    heading("Tổng hợp bốn phát hiện có bằng chứng", 2)
    table(
        "Bảng 8. Phát hiện, bằng chứng và phạm vi diễn giải",
        ["Phát hiện", "Bằng chứng", "Ý nghĩa và giới hạn"],
        [
            (
                "Phân bố giá lệch phải",
                "Median 1.592; mean 2.955,96 THB/đêm; 2.730 ngoại lệ IQR",
                "Báo cáo median cùng IQR để mô tả mức điển hình; chưa xác minh mọi cực trị",
            ),
            (
                "Giá khác nhau theo loại phòng",
                "Shared room median 431,22; Entire home/apt 1.715,44 THB/đêm",
                "Cần so sánh trong phân khúc; không suy ra quan hệ nhân quả",
            ),
            (
                "Liên hệ tuyến tính quy mô–giá nhỏ",
                "Pearson r từ 0,0688 đến 0,1067 trên các cặp hợp lệ",
                "p nhỏ không thay thế độ lớn liên hệ; không dùng EDA toàn mẫu chọn predictor",
            ),
            (
                "Khu vực khác nhau trong cùng loại phòng",
                "Entire home/apt: Bang Phlat 678,03; Samphanthawong 3.115,50 THB/đêm",
                "Chênh lệch mô tả snapshot; chưa tách riêng ảnh hưởng vị trí",
            ),
        ],
        [4, 5.6, 6.4],
    )
    p(
        "Đối với người đọc muốn đối chiếu báo giá, kết quả hỗ trợ lựa chọn nhóm so sánh theo loại phòng và khu vực, đi kèm quy mô mẫu. Báo cáo không đề xuất một mức giá tối ưu, không ước tính doanh thu và không dự báo giá ở thời điểm khác. Phần mô hình bổ trợ cho thấy cần thận trọng với việc đưa ra dự đoán chỉ từ các biến hiện có theo bộ lọc đã chọn."
    )
    heading("KẾT LUẬN")
    p(
        "Đồ án đã xây dựng quy trình kiểm toán, làm sạch và trực quan hóa dữ liệu Airbnb Bangkok từ bản listings ngày 29/06/2026. Sau khi loại 2.082 dòng thiếu giá, 28.987 listing được dùng để trả lời bốn câu hỏi. Giá có trung vị 1.592 THB/đêm, phân phối lệch phải mạnh; trung vị khác nhau theo loại phòng và khu vực, còn liên hệ tuyến tính với các biến quy mô có độ lớn nhỏ. Các bảng và năm hình EDA được sinh từ notebook, có lưu nguồn và số mẫu để tái lập."
    )
    p(
        "Nhánh train/CV không chọn được predictor theo quy tắc hiện tại, nên nhóm chỉ báo cáo mốc trung bình cùng giải trình, chưa xây hồi quy Linear/Ridge cho bộ dữ liệu này. Kết quả không chứng minh dữ liệu không thể mô hình hóa. Giới hạn chính gồm một snapshot niêm yết, thiếu predictor, giá cực trị chưa xác minh đầy đủ, phụ thuộc theo chủ nhà và chưa xác định toàn bộ thuế/phí. Nhóm ưu tiên diễn giải trong phạm vi mẫu quan sát và không khẳng định quan hệ nhân quả hoặc hiệu quả dự báo cho chủ nhà mới."
    )

    page()
    heading("TÀI LIỆU THAM KHẢO")
    references = [
        "[1] Inside Airbnb. Get the Data – Bangkok, detailed listings, bản 29/06/2026. https://insideairbnb.com/get-the-data/. Tệp: https://data.insideairbnb.com/thailand/central-thailand/bangkok/2026-06-29/data/listings.csv.gz. Ngày tải ghi trong manifest: 10/10/2026.",
        "[2] Inside Airbnb. Data Dictionary, phiên bản 4.3 (08/2022). https://docs.google.com/spreadsheets/d/1iWCNJcSutYqpULSQHlNyGInUvHg2BoUGoNRIGa6Szc4/edit. Bản CSV lưu tại docs/sources/inside_airbnb_data_dictionary_v4.3.csv; đối chiếu 11/10/2026. Giấy phép CC BY 4.0: https://creativecommons.org/licenses/by/4.0/.",
        "[3] Phạm Thế Sơn. Bài giảng IE313 – Phân tích và trực quan dữ liệu, UIT. Bài 01–07 và TH 01 trong thư mục slide-study; bản tài liệu môn học, không ghi năm phát hành. Đối chiếu theo trang PDF trong SPECIFICATION.md.",
        "[4] SciPy. scipy.stats.f_oneway; scipy.stats.levene. https://docs.scipy.org/doc/scipy/reference/generated/scipy.stats.f_oneway.html; https://docs.scipy.org/doc/scipy/reference/generated/scipy.stats.levene.html. Truy cập 11/10/2026.",
        "[5] Pandas. pandas.qcut. https://pandas.pydata.org/docs/reference/api/pandas.qcut.html. Truy cập 11/10/2026.",
        "[6] Scikit-learn. r2_score; tài liệu Pipeline và cross-validation. https://scikit-learn.org/stable/modules/generated/sklearn.metrics.r2_score.html. Truy cập 11/10/2026; API cụ thể theo phiên bản trong requirements-lock.txt.",
    ]
    for item in references:
        p(item.split("] ", 1)[1], "TLTK")

    page()
    heading("PHỤ LỤC PHÂN CÔNG NHIỆM VỤ")
    p(
        "Nhóm 10; hạn nộp 18/10/2026. Phân công dưới đây thể hiện trách nhiệm của thành viên và sản phẩm đã có, không thay thế xác nhận kiểm tra chéo của con người. Công cụ AI hỗ trợ triển khai/rà soát; nhóm chịu trách nhiệm nội dung nộp."
    )
    table(
        "Bảng 9. Phân công theo sản phẩm",
        ["Thành viên", "Nhiệm vụ chi tiết", "Sản phẩm / việc cần kiểm tra"],
        [
            (
                "Phạm Quốc Thắng\n25410304",
                "Chốt 4 RQ; EDA và biểu đồ; diễn giải; tổng hợp báo cáo, slide; nộp bài. Kiểm tra chéo dữ liệu và mô hình.",
                "Notebook 02–03, 5 hình EDA, báo cáo Word/PDF. Rà soát diễn giải, nguồn, bố cục và chuẩn bị trình bày.",
            ),
            (
                "Lê Minh Thiện\n25410313",
                "Môi trường; thu thập, kiểm toán và làm sạch; thống kê; train/CV và đánh giá. Kiểm tra chéo hình và bộ nộp.",
                "Notebook 01/04, baseline, thống kê, đánh giá khả thi và kiểm thử. Chạy lại trên máy cá nhân, kiểm tra số liệu và demo.",
            ),
        ],
        [3.5, 6.2, 6.3],
    )
    heading("Danh mục minh chứng tái lập", 2)
    p(
        "01_data_audit_and_cleaning.ipynb: kiểm toán và làm sạch. 02_eda_and_statistical_tests.ipynb: bảng thống kê. 03_visualization_figures.ipynb: hình và catalog. 04_regression_modeling.ipynb: chia tập, chọn biến, mốc trung bình và giải trình."
    )
    p(
        "Dữ liệu nguồn: data/source_manifest.json; từ điển 90 cột: reports/audit/bangkok_columns.csv. Quyết định làm sạch: docs/CLEANING_PLAN.md. Phương pháp: docs/EDA_METHODS.md và docs/MODELING_METHODS.md. Số liệu/hình: reports/stats, reports/figures và reports/modeling. Hướng dẫn thực thi: docs/RUNBOOK.md."
    )
    p(
        "Raw SHA-256: b17b99f9c421c566cc26a1c89a0b3d16b4c7f3dc74e659141900ccbac316b20e. Baseline SHA-256: 961a01ad225e35f989ec81f124925488e392baca2ef3800ff4028b4423d5f9a7."
    )
    p(
        "Mục tiêu nội bộ: nhóm kiểm tra chéo và tổng duyệt trước 17/10/2026. Giờ/kênh nộp, lịch thuyết trình và yêu cầu video được kiểm tra theo thông báo chính thức của giảng viên."
    )
    output = REPORTS / "Nhom10_Bao_cao.docx"
    doc.save(output)
    print(output)


if __name__ == "__main__":
    main()
